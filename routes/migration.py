"""
routes/migration.py — Zagadat Capital Data Migration Portal
============================================================
Dedicated migration blueprint at /migration.
Handles: client accounts, investment portfolios, transaction history, staff users.
"""
from flask import Blueprint, render_template, redirect, url_for, flash, request, send_file
from flask_login import login_required, current_user
from models import (db, ClientAccount, Investment, Transaction, AdminUser,
                    MigrationLog, ADMIN_ROLES)
from routes.admin import permission_required
import io, openpyxl, csv, secrets
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime, date

migration_bp = Blueprint('migration', __name__)

# ── Excel styling (dark gold theme) ───────────────────────────────────────────
HDR_FILL    = PatternFill(start_color='FF1A1A2E', end_color='FF1A1A2E', fill_type='solid')
HDR_FONT    = Font(color='FFC9A84C', bold=True, size=11)
HDR_ALIGN   = Alignment(horizontal='center', vertical='center', wrap_text=True)
BODY_FILL   = PatternFill(start_color='FF16161F', end_color='FF16161F', fill_type='solid')
BODY_FILL2  = PatternFill(start_color='FF1A1A2E', end_color='FF1A1A2E', fill_type='solid')
BODY_FONT   = Font(color='FFF5F5F0', size=10)
THIN_BORDER = Border(
    left=Side(style='thin',   color='FF333355'),
    right=Side(style='thin',  color='FF333355'),
    top=Side(style='thin',    color='FF333355'),
    bottom=Side(style='thin', color='FF333355'),
)


def _style_hdr(ws, row=1):
    for cell in ws[row]:
        cell.fill      = HDR_FILL
        cell.font      = HDR_FONT
        cell.alignment = HDR_ALIGN
        cell.border    = THIN_BORDER


def _style_body(ws, start_row=2):
    for i, row in enumerate(ws.iter_rows(min_row=start_row)):
        fill = BODY_FILL if i % 2 == 0 else BODY_FILL2
        for cell in row:
            cell.fill      = fill
            cell.font      = BODY_FONT
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(vertical='center')


# ── Helpers ───────────────────────────────────────────────────────────────────
def _pd(v):
    """Parse date from various formats."""
    if isinstance(v, (date, datetime)):
        return v.date() if isinstance(v, datetime) else v
    if v:
        for fmt in ['%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y']:
            try:
                return datetime.strptime(str(v).strip(), fmt).date()
            except Exception:
                pass
    return None


def _flt(v, default=0.0):
    try:
        return float(v) if v is not None and str(v).strip() not in ('', 'None') else default
    except Exception:
        return default


def _str(v):
    s = str(v).strip() if v is not None else ''
    return s if s and s.lower() != 'none' else None


def _next_acct_no():
    last = ClientAccount.query.order_by(ClientAccount.id.desc()).first()
    num  = (last.id + 1) if last else 1
    candidate = f'ZC-{num:05d}'
    while ClientAccount.query.filter_by(account_number=candidate).first():
        num += 1
        candidate = f'ZC-{num:05d}'
    return candidate


def _save_log(source, filename, total, ok, fail, errors):
    log = MigrationLog(
        source_system=source,
        filename=filename,
        records_in=total,
        records_ok=ok,
        records_err=fail,
        errors='\n'.join(errors[:100]) if errors else None,
        executed_by=current_user.id,
        status='COMPLETED' if fail == 0 else 'PARTIAL',
    )
    db.session.add(log)
    db.session.commit()


def _read_xlsx_rows(file):
    """Read first sheet of xlsx into list of dicts."""
    wb = openpyxl.load_workbook(file, data_only=True)
    ws = wb.active
    headers = [str(c.value).strip() if c.value else '' for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(v for v in row):
            rows.append(dict(zip(headers, row)))
    return rows


def _read_file_rows(file):
    """Read CSV or XLSX into list of dicts."""
    filename = file.filename
    if filename.lower().endswith('.csv'):
        stream = io.StringIO(file.stream.read().decode('utf-8-sig'))
        return list(csv.DictReader(stream))
    return _read_xlsx_rows(file)


# ─────────────────────────────────────────────────────────────────────────────
# INDEX
# ─────────────────────────────────────────────────────────────────────────────
@migration_bp.route('/')
@login_required
@permission_required('migrate_data')
def index():
    logs = MigrationLog.query.order_by(MigrationLog.executed_at.desc()).limit(30).all()
    return render_template('migration/index.html', logs=logs)


# ─────────────────────────────────────────────────────────────────────────────
# TEMPLATE DOWNLOADS
# ─────────────────────────────────────────────────────────────────────────────
@migration_bp.route('/template/accounts')
@login_required
@permission_required('migrate_data')
def template_accounts():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Client Accounts'

    headers = [
        'account_number', 'full_name', 'account_type', 'phone', 'email',
        'date_of_birth', 'gender', 'marital_status', 'nationality', 'residential_status',
        'address', 'digital_address', 'postal_address', 'city', 'country',
        'id_type', 'id_number', 'ghana_card_number', 'tin', 'csd_number',
        'occupation', 'employer', 'employment_status', 'annual_income_range',
        'source_of_funds', 'investment_objective', 'risk_profile',
        'bank_name', 'bank_branch', 'bank_account_name', 'bank_account_number',
        'beneficiary1_name', 'beneficiary1_relation', 'beneficiary1_pct',
        'opening_balance',
    ]
    ws.append(headers)
    _style_hdr(ws)

    sample = [
        'ZC-00001', 'Kwame Asante Mensah', 'Individual', '0244000001', 'kwame@email.com',
        '1985-03-15', 'Male', 'Married', 'Ghanaian', 'Resident',
        '12 Ring Road East, Accra', 'GA-123-4567', 'P.O. Box AN 1234', 'Accra', 'Ghana',
        'Ghana Card (NIA)', 'GHA-000000001-1', 'GHA-000000001-1', 'C0000000001', 'CSD001',
        'Software Engineer', 'Acme Ltd', 'Employed', 'GHS 60,000 – 150,000',
        'Employment / Salary', 'Capital Growth', 'moderate',
        'GCB Bank', 'Accra Main', 'Kwame Asante Mensah', '1234567890',
        'Ama Mensah', 'Spouse', 50,
        100000,
    ]
    ws.append(sample)
    _style_body(ws)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(
            len(str(col[0].value or '')), 12) + 2

    # Instructions
    ws2 = wb.create_sheet('Instructions')
    ws2['A1'] = 'ZAGADAT CAPITAL — CLIENT ACCOUNTS MIGRATION TEMPLATE'
    ws2['A1'].font = Font(bold=True, size=13)
    ws2.append([''])
    ws2.append(['Column', 'Notes'])
    for col, note in [
        ('account_number', 'Leave blank to auto-generate. Format: ZC-NNNNN'),
        ('account_type', 'Individual | Individual ITF | Joint | Corporate | Pension Fund | NGO'),
        ('date_of_birth', 'YYYY-MM-DD format e.g. 1985-03-15'),
        ('gender', 'Male | Female | Other'),
        ('id_type', 'Ghana Card (NIA) | International Passport | Voter ID | NHIS Card | Driver\'s Licence'),
        ('residential_status', 'Resident | Non-Resident | FATCA Applicable'),
        ('risk_profile', 'conservative | moderate | aggressive'),
        ('investment_objective', 'Capital Preservation | Income Generation | Capital Growth | Balanced'),
        ('beneficiary1_pct', 'Percentage e.g. 50 (for 50%)'),
        ('opening_balance', 'Opening cash balance in GHS. Creates an approved DEPOSIT transaction.'),
    ]:
        ws2.append([col, note])
    ws2.column_dimensions['A'].width = 28
    ws2.column_dimensions['B'].width = 75

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, download_name='zagadat_migration_accounts.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True)


@migration_bp.route('/template/investments')
@login_required
@permission_required('migrate_data')
def template_investments():
    wb = openpyxl.Workbook()

    sheets = {
        'Govt Securities': [
            'account_number', 'security_name', 'issuer', 'sub_type', 'face_value',
            'interest_rate', 'tenor', 'trade_date', 'maturity_date', 'total_cost', 'currency',
        ],
        'Bonds': [
            'account_number', 'security_name', 'issuer', 'isin', 'face_value',
            'coupon_rate', 'coupon_freq', 'clean_price', 'trade_date', 'maturity_date',
            'total_cost', 'rating', 'currency',
        ],
        'Eurobonds': [
            'account_number', 'security_name', 'issuer', 'isin', 'currency',
            'face_value', 'coupon_rate', 'clean_price', 'trade_date', 'maturity_date',
            'total_cost', 'country_of_issue', 'rating',
        ],
        'GSE Equities': [
            'account_number', 'security_name', 'symbol', 'sector',
            'quantity', 'unit_cost', 'current_price', 'total_cost', 'exchange',
        ],
        'Global Equities': [
            'account_number', 'security_name', 'symbol', 'exchange',
            'currency', 'quantity', 'unit_cost', 'current_price', 'total_cost',
        ],
        'Money Market': [
            'account_number', 'security_name', 'issuer', 'sub_type',
            'face_value', 'interest_rate', 'tenor', 'trade_date', 'maturity_date',
            'total_cost', 'currency',
        ],
        'Mutual Funds': [
            'account_number', 'security_name', 'issuer',
            'quantity', 'unit_cost', 'current_price', 'total_cost', 'currency',
        ],
        'Private Equity': [
            'account_number', 'security_name', 'issuer', 'vintage_year',
            'geography', 'stage', 'committed', 'called', 'distributions',
            'nav', 'moic', 'irr', 'currency',
        ],
        'Private Debt': [
            'account_number', 'security_name', 'issuer', 'isin',
            'outstanding', 'coupon_rate', 'trade_date', 'maturity_date',
            'rating', 'pik', 'currency',
        ],
    }

    first = True
    for sheet_name, hdrs in sheets.items():
        if first:
            ws = wb.active
            ws.title = sheet_name
            first = False
        else:
            ws = wb.create_sheet(sheet_name)
        ws.append(hdrs)
        _style_hdr(ws)
        _style_body(ws)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(
                len(str(col[0].value or '')), 12) + 3

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, download_name='zagadat_migration_investments.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True)


@migration_bp.route('/template/transactions')
@login_required
@permission_required('migrate_data')
def template_transactions():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Transactions'
    headers = ['account_number', 'txn_date', 'txn_type', 'description',
               'amount', 'currency', 'reference']
    ws.append(headers)
    _style_hdr(ws)

    samples = [
        ('ZC-00001', '2024-01-15', 'DEPOSIT',    'Initial investment',  100000, 'GHS', 'REF001'),
        ('ZC-00001', '2024-03-01', 'DEPOSIT',    'Top-up Q1 2024',       50000, 'GHS', 'REF002'),
        ('ZC-00001', '2024-06-30', 'DIVIDEND',   'Q2 2024 dividend',      3500, 'GHS', 'DIV001'),
        ('ZC-00001', '2024-09-15', 'COUPON',     '91-day T-Bill coupon',  4200, 'GHS', 'CPN001'),
        ('ZC-00001', '2024-12-01', 'WITHDRAWAL', 'Partial redemption',   20000, 'GHS', 'WDR001'),
    ]
    for s in samples:
        ws.append(s)
    _style_body(ws)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20

    ws2 = wb.create_sheet('Valid TXN Types')
    ws2.append(['TXN Type', 'Description'])
    ws2['A1'].font = HDR_FONT
    ws2['B1'].font = HDR_FONT
    for t, desc in [
        ('DEPOSIT',      'Client deposit / initial investment'),
        ('WITHDRAWAL',   'Client withdrawal / redemption'),
        ('TRANSFER_IN',  'Inbound portfolio transfer'),
        ('TRANSFER_OUT', 'Outbound portfolio transfer'),
        ('DIVIDEND',     'Dividend income received'),
        ('COUPON',       'Coupon / interest payment received'),
        ('FEE',          'Management or performance fee charged'),
        ('BUY',          'Investment purchase'),
        ('SELL',         'Investment sale / redemption'),
    ]:
        ws2.append([t, desc])
    ws2.column_dimensions['A'].width = 18
    ws2.column_dimensions['B'].width = 48

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, download_name='zagadat_migration_transactions.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True)


@migration_bp.route('/template/staff')
@login_required
@permission_required('migrate_data')
def template_staff():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Staff Users'
    headers = ['staff_id', 'full_name', 'email', 'role']
    ws.append(headers)
    _style_hdr(ws)

    samples = [
        ('ZAG-STAFF-001', 'Ama Boateng',  'ama.boateng@zagadatcapital.com',  'ACCOUNT_APPROVER'),
        ('ZAG-STAFF-002', 'Kofi Asante',  'kofi.asante@zagadatcapital.com',  'INVESTMENT_ENTRY'),
        ('ZAG-STAFF-003', 'Efua Mensah',  'efua.mensah@zagadatcapital.com',  'REPORT_VIEWER'),
    ]
    for s in samples:
        ws.append(s)
    _style_body(ws)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(
            len(str(col[0].value or '')), 14) + 4

    ws2 = wb.create_sheet('Valid Roles')
    ws2.append(['Role', 'Level', 'Description'])
    ws2['A1'].font = HDR_FONT
    ws2['B1'].font = HDR_FONT
    ws2['C1'].font = HDR_FONT
    for role, (level, desc) in {
        'SUPER_ADMIN':         (100, 'Full system access — all permissions including migration'),
        'ACCOUNT_SETUP':       (30,  'Create and edit client accounts'),
        'ACCOUNT_APPROVER':    (40,  'Approve / reject KYC accounts'),
        'INVESTMENT_ENTRY':    (50,  'Enter investment transactions'),
        'INVESTMENT_APPROVER': (60,  'Approve investment entries'),
        'REPORT_VIEWER':       (10,  'View reports and dashboards only'),
    }.items():
        ws2.append([role, level, desc])
    ws2.column_dimensions['A'].width = 25
    ws2.column_dimensions['B'].width = 10
    ws2.column_dimensions['C'].width = 55

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, download_name='zagadat_migration_staff.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True)


# ─────────────────────────────────────────────────────────────────────────────
# IMPORT ACCOUNTS
# ─────────────────────────────────────────────────────────────────────────────
@migration_bp.route('/accounts', methods=['POST'])
@login_required
@permission_required('migrate_data')
def import_accounts():
    file = request.files.get('file')
    overwrite = request.form.get('overwrite') == 'on'
    if not file:
        flash('No file uploaded.', 'error')
        return redirect(url_for('migration.index'))

    imported = failed = 0
    errors = []
    filename = file.filename

    try:
        rows = _read_file_rows(file)

        for i, row in enumerate(rows, 2):
            try:
                full_name = _str(row.get('full_name')) or _str(row.get('name'))
                phone     = _str(row.get('phone'))
                acct_no   = _str(row.get('account_number'))

                if not phone:
                    errors.append(f'Row {i}: Missing phone — skipped')
                    failed += 1; continue
                if not full_name:
                    errors.append(f'Row {i}: Missing full_name — skipped')
                    failed += 1; continue

                existing = None
                if acct_no:
                    existing = ClientAccount.query.filter_by(account_number=acct_no).first()
                if not existing:
                    existing = ClientAccount.query.filter_by(phone=phone).first()

                if existing and not overwrite:
                    errors.append(f'Row {i}: Duplicate {phone} — skipped (enable Overwrite to update)')
                    failed += 1; continue

                if existing and overwrite:
                    acc = existing
                else:
                    if not acct_no:
                        acct_no = _next_acct_no()
                    acc = ClientAccount(account_number=acct_no, phone=phone, full_name=full_name)

                # Map all fields
                acc.full_name            = full_name
                acc.phone                = phone
                acc.email                = _str(row.get('email'))               or acc.email
                acc.account_type         = _str(row.get('account_type'))        or acc.account_type or 'Individual'
                acc.gender               = _str(row.get('gender'))              or acc.gender
                acc.marital_status       = _str(row.get('marital_status'))      or acc.marital_status
                acc.nationality          = _str(row.get('nationality'))         or acc.nationality or 'Ghanaian'
                acc.residential_status   = _str(row.get('residential_status')) or acc.residential_status
                acc.address              = _str(row.get('address'))             or acc.address
                acc.digital_address      = _str(row.get('digital_address'))    or acc.digital_address
                acc.postal_address       = _str(row.get('postal_address'))     or acc.postal_address
                acc.city                 = _str(row.get('city'))                or acc.city
                acc.country              = _str(row.get('country'))             or acc.country
                acc.tin                  = _str(row.get('tin'))                 or acc.tin
                acc.csd_number           = _str(row.get('csd_number'))          or acc.csd_number
                acc.id_type              = _str(row.get('id_type'))             or acc.id_type
                acc.id_number            = _str(row.get('id_number'))           or acc.id_number
                acc.ghana_card_number    = _str(row.get('ghana_card_number'))   or acc.ghana_card_number
                acc.occupation           = _str(row.get('occupation'))          or acc.occupation
                acc.employer             = _str(row.get('employer'))            or acc.employer
                acc.employment_status    = _str(row.get('employment_status'))   or acc.employment_status
                acc.annual_income_range  = _str(row.get('annual_income_range')) or acc.annual_income_range
                acc.source_of_funds      = _str(row.get('source_of_funds'))    or acc.source_of_funds
                acc.investment_objective = _str(row.get('investment_objective'))or acc.investment_objective
                acc.risk_profile         = _str(row.get('risk_profile'))        or acc.risk_profile or 'moderate'
                acc.bank_name            = _str(row.get('bank_name'))           or acc.bank_name
                acc.bank_branch          = _str(row.get('bank_branch'))         or acc.bank_branch
                acc.bank_account_name    = _str(row.get('bank_account_name'))   or acc.bank_account_name
                acc.bank_account_number  = _str(row.get('bank_account_number')) or acc.bank_account_number
                acc.beneficiary1_name    = _str(row.get('beneficiary1_name'))   or acc.beneficiary1_name
                acc.beneficiary1_relation= _str(row.get('beneficiary1_relation'))or acc.beneficiary1_relation

                dob = _str(row.get('date_of_birth'))
                if dob:
                    try:
                        acc.date_of_birth = datetime.strptime(dob, '%Y-%m-%d').date()
                    except Exception:
                        pass

                b1p = row.get('beneficiary1_pct')
                if b1p:
                    try:
                        acc.beneficiary1_pct = float(b1p)
                    except Exception:
                        pass

                acc.status      = 'APPROVED'
                acc.approved_by = current_user.id
                acc.approved_at = datetime.utcnow()
                acc.created_by  = acc.created_by or current_user.id

                if not existing:
                    db.session.add(acc)
                db.session.flush()

                # Opening balance → approved DEPOSIT
                opening = _flt(row.get('opening_balance'))
                if opening > 0:
                    txn = Transaction(
                        account_number=acc.account_number,
                        txn_date=date.today(),
                        txn_type='DEPOSIT',
                        description='Opening balance — migration import',
                        amount=opening,
                        amount_ghs=opening,
                        currency='GHS',
                        reference='MIG-OPEN',
                        status='APPROVED',
                        approved_by=current_user.id,
                        approved_at=datetime.utcnow(),
                        created_by=current_user.id,
                    )
                    db.session.add(txn)

                imported += 1

            except Exception as e:
                errors.append(f'Row {i}: {str(e)}')
                failed += 1

        db.session.commit()

    except Exception as e:
        db.session.rollback()
        errors.append(f'File error: {str(e)}')
        failed += 1

    _save_log('Account Migration', filename, imported + failed, imported, failed, errors)
    flash(f'Account migration complete. ✔ Imported: {imported}  ✘ Failed: {failed}',
          'success' if not failed else 'warning')
    return redirect(url_for('migration.index'))


# ─────────────────────────────────────────────────────────────────────────────
# IMPORT INVESTMENTS (portfolio / assets)
# ─────────────────────────────────────────────────────────────────────────────
@migration_bp.route('/investments', methods=['POST'])
@login_required
@permission_required('migrate_data')
def import_investments():
    file = request.files.get('file')
    overwrite = request.form.get('overwrite') == 'on'
    if not file:
        flash('No file uploaded.', 'error')
        return redirect(url_for('migration.index'))

    SHEET_CLASS_MAP = {
        'Govt Securities':  'GOVT_SECURITIES',
        'Bonds':            'BONDS',
        'Eurobonds':        'EUROBONDS',
        'GSE Equities':     'GSE_EQUITIES',
        'Global Equities':  'GLOBAL_EQUITIES',
        'Money Market':     'MONEY_MARKET',
        'Mutual Funds':     'MUTUAL_FUNDS',
        'Private Equity':   'PRIVATE_EQUITY',
        'Private Debt':     'PRIVATE_DEBT',
    }

    imported = failed = 0
    errors   = []
    filename = file.filename

    try:
        wb = openpyxl.load_workbook(file, data_only=True)

        for sheet_name in wb.sheetnames:
            asset_class = SHEET_CLASS_MAP.get(sheet_name)
            if not asset_class:
                continue
            ws      = wb[sheet_name]
            headers = [str(c.value).strip() if c.value else '' for c in ws[1]]

            for i, row_vals in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                rd   = dict(zip(headers, row_vals))
                name = _str(rd.get('security_name')) or _str(rd.get('name'))
                if not name:
                    continue

                acct_no = _str(rd.get('account_number'))
                if not acct_no:
                    errors.append(f'{sheet_name} Row {i}: Missing account_number')
                    failed += 1; continue

                acc = ClientAccount.query.filter_by(account_number=acct_no).first()
                if not acc:
                    errors.append(f'{sheet_name} Row {i}: Account {acct_no} not found')
                    failed += 1; continue

                try:
                    existing = Investment.query.filter_by(
                        account_number=acct_no,
                        security_name=name,
                        asset_class=asset_class,
                    ).first()

                    if existing and not overwrite:
                        errors.append(f'{sheet_name} Row {i}: Duplicate {name} — skipped')
                        failed += 1; continue

                    inv = existing if existing else Investment(
                        account_number=acct_no, asset_class=asset_class)

                    inv.security_name  = name
                    inv.issuer         = _str(rd.get('issuer'))        or inv.issuer
                    inv.symbol         = _str(rd.get('symbol'))         or inv.symbol
                    inv.isin           = _str(rd.get('isin'))           or inv.isin
                    inv.exchange       = _str(rd.get('exchange'))       or inv.exchange
                    inv.sector         = _str(rd.get('sector'))         or inv.sector
                    inv.sub_type       = _str(rd.get('sub_type'))       or inv.sub_type
                    inv.currency       = _str(rd.get('currency'))       or inv.currency or 'GHS'
                    inv.rating         = _str(rd.get('rating'))         or inv.rating
                    inv.country_of_issue = _str(rd.get('country_of_issue')) or inv.country_of_issue
                    inv.geography      = _str(rd.get('geography'))      or inv.geography
                    inv.stage          = _str(rd.get('stage'))          or inv.stage
                    inv.face_value     = _flt(rd.get('face_value'))     or inv.face_value
                    inv.coupon_rate    = _flt(rd.get('coupon_rate'))    or inv.coupon_rate
                    inv.interest_rate  = _flt(rd.get('interest_rate'))  or inv.interest_rate
                    inv.clean_price    = _flt(rd.get('clean_price'))    or inv.clean_price
                    inv.quantity       = _flt(rd.get('quantity'))       or inv.quantity
                    inv.unit_cost      = _flt(rd.get('unit_cost'))      or inv.unit_cost
                    inv.total_cost     = _flt(rd.get('total_cost'))     or inv.total_cost
                    inv.current_price  = _flt(rd.get('current_price'))  or inv.current_price
                    inv.outstanding    = _flt(rd.get('outstanding'))    or inv.outstanding
                    inv.committed      = _flt(rd.get('committed'))      or inv.committed
                    inv.called         = _flt(rd.get('called'))         or inv.called
                    inv.distributions  = _flt(rd.get('distributions'))  or inv.distributions
                    inv.nav            = _flt(rd.get('nav'))            or inv.nav
                    inv.moic           = _flt(rd.get('moic'))           or inv.moic
                    inv.irr            = _flt(rd.get('irr'))            or inv.irr
                    inv.trade_date     = _pd(rd.get('trade_date'))      or inv.trade_date
                    inv.maturity_date  = _pd(rd.get('maturity_date'))   or inv.maturity_date

                    tenor_v = rd.get('tenor')
                    if tenor_v:
                        try:
                            inv.tenor = int(float(str(tenor_v)))
                        except Exception:
                            pass

                    cf_v = rd.get('coupon_freq')
                    if cf_v:
                        try:
                            inv.coupon_freq = int(float(str(cf_v)))
                        except Exception:
                            pass

                    vy_v = rd.get('vintage_year')
                    if vy_v:
                        try:
                            inv.vintage_year = int(float(str(vy_v)))
                        except Exception:
                            pass

                    pik_v = rd.get('pik')
                    if pik_v is not None:
                        inv.pik = str(pik_v).upper() in ('TRUE', '1', 'YES')

                    inv.status      = 'APPROVED'
                    inv.approved_by = current_user.id
                    inv.approved_at = datetime.utcnow()
                    inv.created_by  = inv.created_by or current_user.id

                    if not existing:
                        db.session.add(inv)
                    imported += 1

                except Exception as e:
                    errors.append(f'{sheet_name} Row {i}: {str(e)}')
                    failed += 1

        db.session.commit()

    except Exception as e:
        db.session.rollback()
        errors.append(f'File error: {str(e)}')
        failed += 1

    _save_log('Investment Migration', filename, imported + failed, imported, failed, errors)
    flash(f'Investment migration complete. ✔ Imported: {imported}  ✘ Failed: {failed}',
          'success' if not failed else 'warning')
    return redirect(url_for('migration.index'))


# ─────────────────────────────────────────────────────────────────────────────
# IMPORT TRANSACTIONS
# ─────────────────────────────────────────────────────────────────────────────
@migration_bp.route('/transactions', methods=['POST'])
@login_required
@permission_required('migrate_data')
def import_transactions():
    file = request.files.get('file')
    if not file:
        flash('No file uploaded.', 'error')
        return redirect(url_for('migration.index'))

    VALID_TYPES = {
        'DEPOSIT', 'WITHDRAWAL', 'TRANSFER_IN', 'TRANSFER_OUT',
        'DIVIDEND', 'COUPON', 'FEE', 'BUY', 'SELL',
    }

    imported = failed = 0
    errors   = []
    filename = file.filename

    try:
        rows = _read_file_rows(file)

        for i, row in enumerate(rows, 2):
            try:
                acct_no  = _str(row.get('account_number'))
                txn_type = str(row.get('txn_type') or '').strip().upper()
                amount   = _flt(row.get('amount'))

                if not acct_no:
                    errors.append(f'Row {i}: Missing account_number'); failed += 1; continue
                if txn_type not in VALID_TYPES:
                    errors.append(f"Row {i}: Invalid txn_type '{txn_type}'"); failed += 1; continue
                if amount <= 0:
                    errors.append(f'Row {i}: Invalid amount {amount}'); failed += 1; continue

                acc = ClientAccount.query.filter_by(account_number=acct_no).first()
                if not acc:
                    errors.append(f'Row {i}: Account {acct_no} not found'); failed += 1; continue

                currency = _str(row.get('currency')) or 'GHS'
                txn_date = _pd(row.get('txn_date')) or date.today()

                txn = Transaction(
                    account_number=acct_no,
                    txn_date=txn_date,
                    txn_type=txn_type,
                    description=_str(row.get('description')) or f'{txn_type} — migration import',
                    amount=amount,
                    amount_ghs=amount if currency == 'GHS' else None,
                    currency=currency,
                    reference=_str(row.get('reference')) or 'MIG',
                    status='APPROVED',
                    approved_by=current_user.id,
                    approved_at=datetime.utcnow(),
                    created_by=current_user.id,
                )
                db.session.add(txn)
                imported += 1

            except Exception as e:
                errors.append(f'Row {i}: {str(e)}')
                failed += 1

        db.session.commit()

    except Exception as e:
        db.session.rollback()
        errors.append(f'File error: {str(e)}')
        failed += 1

    _save_log('Transaction Migration', filename, imported + failed, imported, failed, errors)
    flash(f'Transaction migration complete. ✔ Imported: {imported}  ✘ Failed: {failed}',
          'success' if not failed else 'warning')
    return redirect(url_for('migration.index'))


# ─────────────────────────────────────────────────────────────────────────────
# IMPORT STAFF USERS
# ─────────────────────────────────────────────────────────────────────────────
@migration_bp.route('/staff', methods=['POST'])
@login_required
@permission_required('migrate_data')
def import_staff():
    if not current_user.is_super_admin:
        flash('Only the Super Admin can bulk-import staff users.', 'error')
        return redirect(url_for('migration.index'))

    file = request.files.get('file')
    if not file:
        flash('No file uploaded.', 'error')
        return redirect(url_for('migration.index'))

    imported = failed = 0
    errors   = []
    filename = file.filename

    try:
        wb = openpyxl.load_workbook(file, data_only=True)
        ws = wb.active
        headers = [str(c.value).strip().lower() if c.value else '' for c in ws[1]]

        VALID_ROLES = set(ADMIN_ROLES.keys())

        for i, row_vals in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            rd        = dict(zip(headers, row_vals))
            staff_id  = _str(rd.get('staff_id'))
            full_name = _str(rd.get('full_name'))
            email     = _str(rd.get('email'))
            role      = str(rd.get('role') or 'REPORT_VIEWER').strip().upper()

            if not email or not full_name:
                errors.append(f'Row {i}: Missing full_name or email — skipped')
                failed += 1; continue

            if role not in VALID_ROLES:
                role = 'REPORT_VIEWER'

            if AdminUser.query.filter_by(email=email).first():
                errors.append(f'Row {i}: Duplicate email {email} — skipped')
                failed += 1; continue

            if not staff_id:
                count     = AdminUser.query.count() + 1
                staff_id  = f'ZAG-STAFF-{count:04d}'
                while AdminUser.query.filter_by(staff_id=staff_id).first():
                    count += 1
                    staff_id = f'ZAG-STAFF-{count:04d}'
            elif AdminUser.query.filter_by(staff_id=staff_id).first():
                errors.append(f'Row {i}: Duplicate staff_id {staff_id} — skipped')
                failed += 1; continue

            temp_pw = secrets.token_urlsafe(8) + '!'
            user    = AdminUser(
                staff_id=staff_id,
                full_name=full_name,
                email=email,
                role=role,
                created_by=current_user.id,
            )
            user.set_password(temp_pw)
            db.session.add(user)
            db.session.flush()
            imported += 1

        db.session.commit()

    except Exception as e:
        db.session.rollback()
        errors.append(f'File error: {str(e)}')
        failed += 1

    _save_log('Staff Migration', filename, imported + failed, imported, failed, errors)
    flash(f'Staff migration complete. ✔ {imported} users created.  ✘ Failed: {failed}',
          'success' if not failed else 'warning')
    return redirect(url_for('migration.index'))
